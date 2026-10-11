import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6] # Blank slide

# ZeroTouch Executive Palette
C_BG = RGBColor(15, 15, 20)        # Midnight Ink #0F0F14
C_CARD = RGBColor(24, 24, 32)      # Hard Graphite #181820
C_GOLD = RGBColor(212, 175, 106)   # Champagne Gold #D4AF6A
C_GOLD_MUTED = RGBColor(181, 147, 80)
C_WHITE = RGBColor(242, 239, 233)  # Soft White #F2EFE9
C_GRAY = RGBColor(138, 141, 147)   # Cool Gray #8A8D93
C_ACCENT = RGBColor(74, 111, 165)  # Slate Blue #4A6FA5
C_GREEN = RGBColor(46, 204, 113)   # Feedback Emerald

def create_base_slide():
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()
    return slide

def add_header(slide, title_text, category_text="ZEROTOUCH OS • UAT HACKATHON"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = C_GOLD
    p_c.font.name = 'Segoe UI'

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.9))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE
    p_t.font.name = 'Georgia'

def add_card(slide, left, top, width, height, title, body_bullets, accent_color=C_GOLD, live_badge=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = accent_color
    card.line.width = Pt(1)

    pad = Inches(0.25)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_head = tf.paragraphs[0]
    p_head.text = title
    p_head.font.size = Pt(18)
    p_head.font.bold = True
    p_head.font.color.rgb = accent_color
    p_head.font.name = 'Segoe UI'

    if live_badge:
        p_badge = tf.add_paragraph()
        p_badge.text = f"[ LIVE DEMO ACTION: {live_badge} ]"
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = C_GREEN
        p_badge.font.name = 'Consolas'

    for b in body_bullets:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_WHITE
        p.font.name = 'Segoe UI'
        p.space_before = Pt(6)

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
s1 = create_base_slide()
if os.path.exists('zerotouch.png'):
    s1.shapes.add_picture('zerotouch.png', Inches(5.86), Inches(1.3), Inches(1.6), Inches(1.6))

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(11.33), Inches(3.5))
tf1 = tb1.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.alignment = PP_ALIGN.CENTER
p1.text = "ZeroTouch OS"
p1.font.size = Pt(54)
p1.font.bold = True
p1.font.color.rgb = C_GOLD
p1.font.name = 'Georgia'

p2 = tf1.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.text = "THE INVISIBLE SPATIAL PERIPHERAL FOR WINDOWS"
p2.font.size = Pt(15)
p2.font.bold = True
p2.font.color.rgb = C_WHITE
p2.font.name = 'Consolas'
p2.space_before = Pt(10)

p3 = tf1.add_paragraph()
p3.alignment = PP_ALIGN.CENTER
p3.text = "Executive Computer Vision • 100% Local Inferencing • Zero False Positives"
p3.font.size = Pt(14)
p3.font.color.rgb = C_GRAY
p3.font.name = 'Segoe UI'
p3.space_before = Pt(6)

p4 = tf1.add_paragraph()
p4.alignment = PP_ALIGN.CENTER
p4.text = "UAT Hackathon 2026: See / Analyze / Build"
p4.font.size = Pt(13)
p4.font.color.rgb = C_GOLD_MUTED
p4.font.name = 'Segoe UI'
p4.space_before = Pt(28)

# ==========================================
# SLIDE 2: The Core Problem & Thesis
# ==========================================
s2 = create_base_slide()
add_header(s2, "The Spatial Computing Problem: Why Cameras Stay Off", "01 • THE PROBLEM & THESIS")
add_card(s2, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8), 
         "The Problem with Traditional CV",
         ["Accidental Triggers: Scratching face, reaching for coffee, or typing fires unintended commands.",
          "Privacy Paranoia: Users distrust cameras streaming frames into cloud servers or storing them on disk.",
          "Fatigue & Distance Limits: Standing up or stepping back to present instantly breaks hand tracking.",
          "Hardware Lockouts: Most CV software locks out Zoom, Teams, or Google Meet from accessing the webcam."],
         C_GOLD)

add_card(s2, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8),
         "The ZeroTouch Philosophy",
         ["The Invisible Peripheral: Stays silent in the background until deliberate intent is mathematically verified.",
          "100% Volatile RAM: Zero frames stored, zero network requests, zero telemetry leakage.",
          "Multimodal Intent Verification: Real physical typing detection, body ownership reach bounds, and gaze yaw tracking.",
          "Native OS Integration: Leverages Windows Camera Frame Server so Zoom, Meet, and ZeroTouch run simultaneously."],
         C_WHITE)

# ==========================================
# SLIDE 3: Context-Aware Intelligence
# ==========================================
s3 = create_base_slide()
add_header(s3, "Context-Aware OS Control: One Engine, Three Modes", "02 • SPATIAL WORKFLOW ENGINE")
add_card(s3, Inches(0.8), Inches(1.9), Inches(3.7), Inches(4.8),
         "Presentation Mode",
         ["Targets: PowerPoint, Google Slides, Acrobat PDF.",
          "Back-of-Hand Swipe: Instant Next / Previous Slide advance.",
          "Open Palm Hold (0.5s): Toggles Blank/Resume presentation screen (B key).",
          "Stage-Optimized: Gaze lockouts bypassed so presenters can look at slides or audience."],
         C_GOLD,
         "Swipe back of hand to advance this slide right now!")

add_card(s3, Inches(4.8), Inches(1.9), Inches(3.7), Inches(4.8),
         "Video Meeting Mode",
         ["Targets: Google Meet, Zoom, MS Teams, Webex.",
          "Ghost-Mute (Shh): Index finger to mouth toggles hardware microphone mute directly via Windows Core Audio.",
          "Palm Hold: Instant video camera toggle (Alt + V).",
          "Swipe: Toggle participant / gallery views."],
         C_WHITE,
         "Place finger on lips to trigger hardware mic mute!")

add_card(s3, Inches(8.8), Inches(1.9), Inches(3.7), Inches(4.8),
         "Media & Default Mode",
         ["Targets: Spotify, YouTube, Netflix, Desktop.",
          "Palm Hold: Smooth Media Play/Pause.",
          "Back-of-Hand Swipe: Next/Prev audio track.",
          "Dynamic Vertical Pinch: Continuous Master Volume slide with live floating capsule feedback."],
         C_ACCENT,
         "Pinch thumb and index finger to slide volume!")

# ==========================================
# SLIDE 4: Computer Vision Defense Architecture
# ==========================================
s4 = create_base_slide()
add_header(s4, "Noise Rejection: How We Eliminate False Positives", "03 • COMPUTER VISION DEFENSE")
add_card(s4, Inches(0.8), Inches(1.9), Inches(5.6), Inches(2.3),
         "Smart Activity Lockout (Type and Click)",
         ["Low-level background hooks (Pynput) detect physical keyboard and mouse activity.",
          "While typing or resting hands on the keyboard, gestures sleep for 1.8s."],
         C_GOLD,
         "Type on keyboard -> Gestures instantly sleep")

add_card(s4, Inches(0.8), Inches(4.45), Inches(5.6), Inches(2.3),
         "Gaze-Clutch (Look-to-Control)",
         ["MediaPipe keypoints track head yaw and pupillary distance in real time.",
          "Looking away at a second monitor immediately clutches gestures to standby."],
         C_WHITE,
         "Turn head away -> HUD shows GAZE: AWAY")

add_card(s4, Inches(6.8), Inches(1.9), Inches(5.7), Inches(2.3),
         "ShoulderGuard Intruder Shield",
         ["BlazeFace scans the workspace for shoulder surfers.",
          "If a second face appears behind you, fires Win+D to hide sensitive work instantly."],
         C_ACCENT,
         "Second person peers over shoulder -> Desktop minimizes")

add_card(s4, Inches(6.8), Inches(4.45), Inches(5.7), Inches(2.3),
         "Body Ownership & Anti-Jitter Physics",
         ["Calculates Euclidean reach boundaries between hand wrist and face center.",
          "Passersby or background people cannot steal control; hallucinated joints freeze."],
         C_GOLD,
         "Hands outside reach radius are rejected")

# ==========================================
# SLIDE 5: Stage Mode & Camera Sharing
# ==========================================
s5 = create_base_slide()
add_header(s5, "Engineering Breakthroughs: Stage Range & Camera Sharing", "04 • KEY INNOVATIONS")
add_card(s5, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8),
         "Stage & Distance Mode",
         ["The Challenge: Webcams at 6-10 feet compress face resolution down to 10 pixels, causing false look away lockouts.",
          "Dynamic Distance Ratio: Real-time calculation of face width relative to screen width (fw / w).",
          "Automatic Focus Bypass: Gaze lockout gracefully disengages when presenting or at distance.",
          "6.0x Arm Reach: Expands presenter reach envelope so wide, expressive stage gestures register smoothly."],
         C_GOLD,
         "Step back 8 feet and wave to advance slides!")

add_card(s5, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8),
         "Multi-App Camera Sharing",
         ["The Challenge: Windows UVC camera drivers enforce exclusive hardware locks, crashing Zoom or Google Meet if CV is running.",
          "Native Frame Server: Configured OpenCV Media Foundation (CAP_MSMF) to hook svchost.exe -k Camera.",
          "Zero Virtual Camera Latency: No third-party drivers or OBS needed.",
          "1-Click Dashboard Setup: Integrated OS deep-link (ms-settings:camera) with in-app guidance for instant multi-app streaming."],
         C_WHITE,
         "Google Meet and ZeroTouch run on same webcam!")

# ==========================================
# SLIDE 6: System Architecture
# ==========================================
s6 = create_base_slide()
add_header(s6, "System Architecture: High-Throughput & Low Footprint", "05 • TECHNICAL ARCHITECTURE")
add_card(s6, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8),
         "High-Throughput IPC Pipeline",
         ["Shared Memory (multiprocessing.shared_memory): Zero-copy 30 FPS telemetry and downscaled feed streamed between engine and GUI.",
          "Ultra-Lightweight Dashboard: CustomTkinter widget (350x600) with collapsible video preview to save CPU cycles.",
          "Windows System Tray: Completely hideable to system tray with persistent background execution.",
          "Floating Capsule HUD: Tkinter topmost micro-pill that drops down for 1.2s to confirm gestures without blocking work."],
         C_ACCENT)

add_card(s6, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8),
         "Tech Stack & Standards",
         ["Computer Vision: Google MediaPipe Tasks API (HandLandmarker + BlazeFace)",
          "Frame Ingestion: OpenCV 5.0 with MSMF Frame Server Broker",
          "Hardware Audio Hook: Windows Core Audio via Pycaw (Endpoint Volume API)",
          "OS Automation: PyAutoGUI, Pynput, PyWin32 (Win32 API)",
          "Acoustic Haptics: Windows Sound API (Winsound) with custom wave chimes",
          "Packaging: Native 1-Click Batch Launcher (Start_Aero.bat) with virtualenv isolation"],
         C_GOLD)

# ==========================================
# SLIDE 7: Live Demo Guide & Conclusion
# ==========================================
s7 = create_base_slide()
add_header(s7, "Live Demonstration: The 4 Golden Demos", "06 • LIVE DEMONSTRATION & WRAP-UP")
add_card(s7, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8),
         "4 Interactive Demonstrations for Judges",
         ["1. Slide Deck Navigation: Back-of-hand swipe to cycle presentation slides forward and backward right now.",
          "2. Ghost-Mute: Index finger to mouth to toggle Windows hardware microphone mute with haptic chime.",
          "3. Smart Activity Lockout: Typing on the keyboard instantly freezes gestures to prevent accidental input.",
          "4. Stage Distance Reach: Stepping back across the room to show seamless long-range slide control."],
         C_GREEN,
         "Ready for Live Judges Hands-On")

add_card(s7, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8),
         "Summary & Future Vision",
         ["Summary: ZeroTouch transforms standard webcams into executive, enterprise-grade spatial interfaces without cloud risk or false inputs.",
          "Future Roadmap:",
          "  • Multimodal head-tilt micro gestures for accessibility.",
          "  • Native Windows 11 Copilot / Taskbar widget integration.",
          "  • Hardware acceleration optimizations via ONNX Runtime / DirectML.",
          "GitHub: https://github.com/Philphil-org/UAT-Hackathon"],
         C_GOLD)

output_path = 'ZeroTouch_Presentation.pptx'
prs.save(output_path)
print(f'Successfully generated {output_path}')
